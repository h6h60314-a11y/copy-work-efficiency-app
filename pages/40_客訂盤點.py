#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""選擇客訂（可多檔）、庫存與儲位 Excel，自動產生客訂盤點明細。"""
from __future__ import annotations
import io
from datetime import datetime, date

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from common_ui import card_close, card_open, inject_logistics_theme, set_page

HEADERS = ['商品號','實際數量','實際效期','儲位','國際條碼','條碼2','全名','數量','已試算','已揀取','Canuseqty','商品效期','棚別']


def norm(v):
    if v is None or pd.isna(v):
        return ''
    s = str(v).strip()
    return s[:-2] if s.endswith('.0') and s[:-2].isdigit() else s


def load_excel(source):
    name = getattr(source, 'name', str(source))
    suffix = name.lower().rsplit('.', 1)[-1] if '.' in name else ''
    if suffix not in ('xls', 'xlsx'):
        raise ValueError(f'不支援的格式：{name}，請使用 .xls/.xlsx')
    engine = 'xlrd' if suffix == 'xls' else 'openpyxl'
    if hasattr(source, 'seek'):
        source.seek(0)
    try:
        frame = pd.read_excel(source, dtype=str, engine=engine, keep_default_na=False)
    except ImportError as ex:
        raise RuntimeError('缺少 Excel 讀取套件，請執行：pip install pandas openpyxl xlrd') from ex
    frame.columns = [str(c).strip() for c in frame.columns]
    return frame


def make_records(orders, stock, location):
    for idx, order in enumerate(orders, 1):
        if '商品' not in order.columns:
            raise ValueError(f'客訂第 {idx} 份缺少「商品」欄位')
    expected = {'商品號','儲位','國際條碼','條碼2','全名','數量','已試算','已揀取','Canuseqty','商品效期'}
    missing = expected - set(stock.columns)
    if missing:
        raise ValueError('庫存缺少欄位：' + '、'.join(sorted(missing)))
    if '儲位' not in location.columns or '棚別' not in location.columns:
        raise ValueError('儲位明細缺少「儲位」或「棚別」欄位')
    products = {norm(value) for df in orders for value in df['商品'] if norm(value)}
    locmap = {}
    for _, r in location.iterrows():
        code = norm(r['儲位'])
        if code and code not in locmap:
            locmap[code] = norm(r['棚別'])
    rows = []
    bad_locations = set()
    for _, r in stock.iterrows():
        product = norm(r['商品號'])
        if product not in products:
            continue
        slot = norm(r['儲位'])
        shelf = locmap.get(slot, '')
        if not shelf:
            bad_locations.add(slot)
        rows.append([product, '', '', slot, norm(r['國際條碼']), norm(r['條碼2']), norm(r['全名']),
                     norm(r['數量']), norm(r['已試算']), norm(r['已揀取']), norm(r['Canuseqty']),
                     norm(r['商品效期']), shelf])
    seen = {r[0] for r in rows}
    info = {'商品數': len(products), '盤點筆數': len(rows),
            '無庫存商品': sorted(products - seen), '無棚別儲位': sorted(bad_locations)}
    return rows, info


def as_date(value):
    if isinstance(value, (date, datetime)):
        return value
    s = norm(value)
    if not s:
        return None
    try:
        return datetime.fromisoformat(s)
    except ValueError:
        return s


def as_number(value):
    s = norm(value)
    if not s:
        return None
    try:
        n = float(s.replace(',', ''))
        return int(n) if n.is_integer() else n
    except ValueError:
        return s


def save_report(rows, info, output):
    wb = Workbook()
    ws = wb.active
    ws.title = '盤點明細'
    ws.append(HEADERS)
    for row in rows:
        transformed = [as_number(v) if i in (7,8,9,10) else as_date(v) if i == 11 else v
                       for i, v in enumerate(row)]
        ws.append(transformed)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:M{max(ws.max_row, 2)}'
    for cell in ws[1]:
        cell.font = Font(name='Microsoft JhengHei', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='234B70')
        cell.alignment = Alignment(horizontal='center')
    widths = [16,14,15,17,20,18,50,12,12,12,15,17,16]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cells in ws.iter_rows(min_row=2):
        for c in cells:
            c.font = Font(name='Microsoft JhengHei', size=10)
            if c.column in (2,3):
                c.fill = PatternFill('solid', fgColor='FFF2CE')
            if c.column == 12 and isinstance(c.value, (date, datetime)):
                c.number_format = 'yyyy/mm/dd'
    if info['無庫存商品'] or info['無棚別儲位']:
        w = wb.create_sheet('待確認')
        w.append(['類型','代碼','說明'])
        for item in info['無庫存商品']:
            w.append(['無庫存商品', item, '客訂有此商品，但庫存檔沒有庫存紀錄'])
        for item in info['無棚別儲位']:
            w.append(['無棚別儲位', item, '儲位明細無法對應棚別'])
        for k, size in [('A', 20), ('B', 25), ('C', 45)]:
            w.column_dimensions[k].width = size
    wb.save(output)
    return output


def execute(order_paths, stock_path, location_path, output_path):
    if not order_paths:
        raise ValueError('至少需要一份客訂明細')
    orders = [load_excel(p) for p in order_paths]
    stock = load_excel(stock_path)
    locations = load_excel(location_path)
    rows, info = make_records(orders, stock, locations)
    if not rows:
        raise ValueError('找不到客訂商品對應的庫存紀錄，請確認來源檔案')
    save_report(rows, info, output_path)
    return info


def build_report(order_files, stock_file, location_file):
    orders = [load_excel(file) for file in order_files]
    stock = load_excel(stock_file)
    locations = load_excel(location_file)
    rows, info = make_records(orders, stock, locations)
    if not rows:
        raise ValueError('找不到客訂商品對應的庫存紀錄，請確認來源檔案')

    output = io.BytesIO()
    save_report(rows, info, output)
    output.seek(0)
    return output.getvalue(), info


def main():
    inject_logistics_theme()
    set_page('客訂盤點', icon='🔎', subtitle='上傳客訂、庫存與儲位明細，自動產生盤點 Excel。')

    card_open('🔎 客訂盤點明細產生')
    st.caption('支援 .xlsx 與 .xls；客訂未出明細可一次上傳多份。')

    order_files = st.file_uploader(
        '① 客訂未出明細（可多選）',
        type=['xlsx', 'xls'],
        accept_multiple_files=True,
    )
    stock_file = st.file_uploader('② 庫存明細', type=['xlsx', 'xls'])
    location_file = st.file_uploader('③ 儲位明細', type=['xlsx', 'xls'])

    ready = bool(order_files and stock_file and location_file)
    if st.button('開始產生盤點明細', type='primary', use_container_width=True, disabled=not ready):
        try:
            report, info = build_report(order_files, stock_file, location_file)
            st.session_state['custom_order_inventory_report'] = report
            st.session_state['custom_order_inventory_info'] = info
        except Exception as ex:
            st.session_state.pop('custom_order_inventory_report', None)
            st.session_state.pop('custom_order_inventory_info', None)
            st.error(f'處理失敗：{ex}')

    report = st.session_state.get('custom_order_inventory_report')
    info = st.session_state.get('custom_order_inventory_info')
    if report and info:
        st.success('客訂盤點明細已完成。')
        col1, col2, col3, col4 = st.columns(4)
        col1.metric('客訂商品', info['商品數'])
        col2.metric('盤點筆數', info['盤點筆數'])
        col3.metric('無庫存商品', len(info['無庫存商品']))
        col4.metric('缺棚別儲位', len(info['無棚別儲位']))

        st.download_button(
            '下載客訂盤點 Excel',
            data=report,
            file_name=f'客訂盤點_{datetime.now():%Y%m%d_%H%M}.xlsx',
            mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            use_container_width=True,
        )
    elif not ready:
        st.info('請依序上傳客訂未出明細、庫存明細與儲位明細。')

    card_close()


if __name__ == '__main__':
    main()
